from pathlib import Path
from PIL import Image, ImageOps
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\ASUS Zenbook\Documents\ChatGPT\Pinballl")
OUT = ROOT / "Ho-so-thuyet-minh-may-ban-bi-Pinball-trinh-TTTM.docx"
ASSET_DIR = ROOT / "dossier_assets"
ASSET_DIR.mkdir(exist_ok=True)

SOURCE_IMAGES = {
    "front": Path(r"C:\Trang Nguyễn\Starthing\Bumbum\máy bán hàng +\Bắn bi\ảnh\IMG_2353.JPG"),
    "playfield": Path(r"C:\Trang Nguyễn\Starthing\Bumbum\máy bán hàng +\Bắn bi\ảnh\IMG_2357.JPG"),
    "controls": Path(r"C:\Trang Nguyễn\Starthing\Bumbum\máy bán hàng +\Bắn bi\ảnh\IMG_2359.JPG"),
    "installation": Path(r"C:\Users\ASUS Zenbook\AppData\Local\Temp\codex-clipboard-ab4aeb0e-bd17-47e6-a8cb-4ff6fb8020ce.png"),
    "dimensions": Path(r"C:\Users\ASUS Zenbook\AppData\Local\Temp\codex-clipboard-cd321c2d-1684-45e5-869f-c61e46885911.png"),
}


def prepare_image(src: Path, name: str, max_side=2200, white_bg=False):
    dst = ASSET_DIR / f"{name}.jpg"
    with Image.open(src) as im:
        im = ImageOps.exif_transpose(im).convert("RGB")
        if white_bg:
            bg = Image.new("RGB", im.size, "white")
            bg.paste(im)
            im = bg
        im.thumbnail((max_side, max_side), Image.Resampling.LANCZOS)
        im.save(dst, quality=92, optimize=True)
    return dst


IMAGES = {key: prepare_image(path, key) for key, path in SOURCE_IMAGES.items()}


def set_run_font(run, name="Times New Roman", size=13, bold=None, italic=None, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9D9D9", size="6"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def remove_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = tblPr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        borders.append(node)


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


def set_repeat_table_header_if_needed(table):
    if table.rows:
        set_repeat_table_header(table.rows[0])


def set_keep_with_next(paragraph, value=True):
    paragraph.paragraph_format.keep_with_next = value


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, text, fld_end])
    set_run_font(run, size=13)


def add_text(paragraph, text, bold=False, italic=False, size=13, color="000000"):
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic, color=color)
    return run


def add_body(text="", indent=True, bold_lead=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)
    if indent:
        p.paragraph_format.first_line_indent = Cm(1.0)
    if bold_lead and text.startswith(bold_lead):
        add_text(p, bold_lead, bold=True)
        add_text(p, text[len(bold_lead):])
    else:
        add_text(p, text)
    return p


def add_bullet(text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.75 + level * 0.6)
    p.paragraph_format.first_line_indent = Cm(-0.35)
    p.paragraph_format.line_spacing = 1.12
    p.paragraph_format.space_after = Pt(3)
    add_text(p, text)
    return p


def add_heading(text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(10 if level == 1 else 7)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.05
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return p


def add_table(headers, rows, widths=None, font_size=11.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, "1F4E78")
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        add_text(p, h, bold=True, size=font_size, color="FFFFFF")
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_border(cells[i])
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ridx % 2 == 1:
                set_cell_shading(cells[i], "EEF5FA")
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 and len(headers) > 2 else WD_ALIGN_PARAGRAPH.LEFT
            add_text(p, str(value), size=font_size)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    set_repeat_table_header_if_needed(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_figure(path, caption, width_cm=15.0, page_break_before=False):
    if page_break_before:
        doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(7)
    add_text(c, caption, italic=True, size=11)
    return p


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21.0)
section.page_height = Cm(29.7)
section.top_margin = Cm(2.0)
section.bottom_margin = Cm(2.0)
section.left_margin = Cm(3.0)
section.right_margin = Cm(2.0)
section.header_distance = Cm(0.8)
section.footer_distance = Cm(0.8)
section.different_first_page_header_footer = True
add_page_number(section.header.paragraphs[0])

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Times New Roman"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
normal.font.size = Pt(13)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15

title_style = styles["Title"]
title_style.font.name = "Times New Roman"
title_style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
title_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
title_style.font.size = Pt(20)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)

for level, size in [(1, 15), (2, 13)]:
    s = styles[f"Heading {level}"]
    s.font.name = "Times New Roman"
    s._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    s._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    s._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = RGBColor(0, 0, 0)

# Cover header
hdr = doc.add_table(rows=1, cols=2)
hdr.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr.autofit = False
remove_table_borders(hdr)
hdr.columns[0].width = Cm(7.5)
hdr.columns[1].width = Cm(8.5)
left, right = hdr.rows[0].cells
for cell in (left, right):
    set_cell_margins(cell, 0, 40, 0, 40)

p = left.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_text(p, "CÔNG TY TNHH CÔNG NGHỆ\nVIỆT CORE VIỆT NAM", bold=True, size=12)
p2 = left.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_text(p2, "Số: ……/2026/HS-VCC", size=12)

p = right.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_text(p, "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", bold=True, size=12)
p2 = right.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_text(p2, "Độc lập - Tự do - Hạnh phúc", bold=True, size=13)
p3 = right.add_paragraph()
p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_text(p3, "TP. Hồ Chí Minh, ngày …… tháng …… năm 2026", italic=True, size=11.5)

doc.add_paragraph().paragraph_format.space_after = Pt(12)
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(5)
add_text(p, "HỒ SƠ THUYẾT MINH THIẾT BỊ", bold=True, size=20)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(5)
add_text(p, "MÁY BẮN BI PINBALL ĐỔI QUÀ", bold=True, size=18)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(12)
add_text(p, "Phục vụ xem xét lắp đặt và vận hành tại trung tâm thương mại", italic=True, size=13)

add_figure(IMAGES["installation"], "Hình ảnh thực tế máy bắn bi Pinball tại điểm trưng bày", width_cm=10.8)

meta = doc.add_table(rows=4, cols=2)
meta.alignment = WD_TABLE_ALIGNMENT.CENTER
meta.autofit = False
remove_table_borders(meta)
meta_data = [
    ("Đơn vị cung cấp", "CÔNG TY TNHH CÔNG NGHỆ VIỆT CORE VIỆT NAM"),
    ("Mã số thuế", "0319163023"),
    ("Địa chỉ", "24 đường số 11, CityLand ParkHill, Gò Vấp, TP. Hồ Chí Minh"),
    ("Liên hệ", "037 599 3345 | tnhhcongnghevietcore@gmail.com"),
]
for row, (a, b) in zip(meta.rows, meta_data):
    set_cell_margins(row.cells[0], 45, 80, 45, 80)
    set_cell_margins(row.cells[1], 45, 80, 45, 80)
    add_text(row.cells[0].paragraphs[0], a + ":", bold=True, size=11.5)
    add_text(row.cells[1].paragraphs[0], b, size=11.5)

doc.add_page_break()

# Introduction and document status
add_heading("1. Mục đích và phạm vi hồ sơ", 1)
add_body("Công ty TNHH Công nghệ Việt Core Việt Nam kính trình Ban Quản lý Trung tâm thương mại hồ sơ thuyết minh máy bắn bi Pinball đổi quà để xem xét phương án tiếp nhận, bố trí, kết nối kỹ thuật và vận hành thiết bị tại khu vực được phê duyệt.")
add_body("Hồ sơ tập trung vào nhận diện sản phẩm, thông số kỹ thuật, nguyên lý hoạt động, phương thức thanh toán và trả quà, phần mềm quản lý, yêu cầu hạ tầng, biện pháp an toàn, hồ sơ nguồn gốc và hình ảnh đối chiếu thực tế.")

add_heading("2. Cơ sở trình bày và lưu ý áp dụng", 1)
add_body("Bố cục nội dung được sắp xếp theo nhóm thông tin quản lý đối với máy trò chơi điện tử từng được đề cập tại Thông tư số 03/1998/TT-BVHTT ngày 22/6/1998. Thông tư này đã được Thông tư số 08/2000/TT-BVHTT thay thế từ ngày 01/6/2000; vì vậy hồ sơ không sử dụng Thông tư 03/1998/TT-BVHTT như căn cứ hiệu lực hiện hành.")
add_body("Thể thức trình bày văn bản áp dụng khổ A4, phông Times New Roman và định lề theo Phụ lục I Nghị định số 30/2020/NĐ-CP ngày 05/3/2020 về công tác văn thư. Việc chấp thuận lắp đặt, thành phần chứng từ và yêu cầu kỹ thuật cuối cùng thực hiện theo quy định hiện hành, nội quy của trung tâm thương mại và yêu cầu của cơ quan quản lý tại địa phương.")

add_heading("3. Danh mục nội dung hồ sơ", 1)
for item in [
    "Thông tin nhận diện thiết bị và đơn vị cung cấp.",
    "Mô tả chức năng, nguyên lý chơi và cơ chế đổi quà.",
    "Thông số kỹ thuật và yêu cầu hạ tầng lắp đặt.",
    "Hệ thống thanh toán và phần mềm quản lý vận hành.",
    "Phương án bố trí, vận hành, an toàn và xử lý sự cố.",
    "Tình trạng hồ sơ pháp lý, chứng từ và danh mục cần đính kèm.",
    "Ảnh mô phỏng kích thước và ảnh thực tế của thiết bị.",
]:
    add_bullet(item)

doc.add_page_break()

add_heading("4. Thông tin nhận diện thiết bị", 1)
add_table(["Hạng mục", "Thông tin"], [
    ("Tên thiết bị", "Máy bắn bi Pinball đổi quà"),
    ("Tên hiển thị trên máy", "BONBO PINBALL"),
    ("Loại hình", "Thiết bị trò chơi giải trí điện tử có tích lũy điểm và đổi quà hiện vật"),
    ("Xuất xứ", "Trung Quốc; hàng nhập khẩu chính ngạch theo hồ sơ do nhà cung cấp xuất trình"),
    ("Đơn vị cung cấp", "Công ty TNHH Công nghệ Việt Core Việt Nam"),
    ("Đối tượng sử dụng", "Khách hàng tại khu vui chơi, khu giải trí và khu vực dịch vụ trong trung tâm thương mại"),
    ("Số lượng đề xuất", "…… máy"),
    ("Vị trí đề xuất", "Tầng/khu vực: …………………………………………………………………"),
    ("Mã model / số seri", "Bổ sung theo tem nhãn của từng máy trước khi bàn giao"),
], widths=[5.0, 10.7], font_size=11.5)

add_heading("5. Mô tả sản phẩm và mục đích sử dụng", 1)
add_body("Máy bắn bi Pinball là thiết bị vui chơi giải trí kết hợp thao tác bắn bi, ghi nhận điểm và đổi quà. Người chơi thanh toán để nhận lượt chơi, sử dụng cần bắn bi, theo dõi kết quả trên bộ hiển thị điểm và dùng tổng điểm đạt được để nhận quà theo bảng điểm được công bố trên máy.")
add_body("Thiết bị có chín ngăn trưng bày quà. Mỗi ngăn được gắn mức điểm tương ứng; khi đủ điểm, người chơi lựa chọn ngăn quà và hệ thống mở ngăn để nhận quà. Mô hình vận hành tự phục vụ giúp giảm thao tác thu tiền và trả quà trực tiếp của nhân viên.")
add_body("Nguyên tắc vận hành đề xuất tại trung tâm thương mại: điểm thưởng chỉ dùng đổi quà hiện vật theo bảng điểm; không quy đổi thành tiền mặt; không tổ chức đặt cược; nội dung và hình ảnh trên máy phù hợp hoạt động vui chơi giải trí lành mạnh.")

doc.add_page_break()

add_heading("6. Thông số kỹ thuật", 1)
add_table(["STT", "Thông số", "Giá trị / mô tả"], [
    ("1", "Kích thước thân máy", "Rộng 510 mm x sâu 778 mm x cao 1.700 mm"),
    ("2", "Diện tích chiếm chỗ thân máy", "Khoảng 0,40 m²"),
    ("3", "Diện tích vận hành khuyến nghị", "Khoảng 1 m²/máy, bao gồm ghế và vùng thao tác; chưa bao gồm lối đi chung theo yêu cầu TTTM"),
    ("4", "Trọng lượng", "Khoảng 60 kg/máy"),
    ("5", "Tải trọng phân bố tham khảo", "Khoảng 151 kg/m² trên diện tích đáy máy; chưa bao gồm người chơi và ghế"),
    ("6", "Nguồn điện", "220 V AC"),
    ("7", "Công suất tiêu thụ", "100-120 W/máy"),
    ("8", "Vật liệu thân máy", "Nhựa và các cụm linh kiện điện - điện tử; đặc tính chống cháy cần đối chiếu chứng từ nhà sản xuất nếu TTTM yêu cầu"),
    ("9", "Số ngăn quà", "09 ngăn"),
    ("10", "Hiển thị", "Bảng điểm điện tử và đèn LED"),
    ("11", "Kết nối", "Có kết nối mạng phục vụ thanh toán QR và quản lý từ xa"),
    ("12", "Phương thức thanh toán", "Tiền mặt và chuyển khoản bằng mã QR"),
    ("13", "Phụ kiện kèm theo", "01 ghế ngồi chơi/máy"),
    ("14", "Điều kiện môi trường", "Lắp đặt trong nhà, khu vực khô ráo, thông thoáng, tránh nước và nguồn nhiệt trực tiếp"),
], widths=[1.2, 5.0, 9.5], font_size=10.8)

add_heading("7. Nhu cầu hạ tầng tại điểm đặt", 1)
add_table(["Hạng mục", "Yêu cầu đề xuất"], [
    ("Nguồn điện", "01 ổ cắm 220 V/máy; hệ thống ổ cắm, tiếp địa và bảo vệ rò điện tuân theo tiêu chuẩn kỹ thuật của TTTM."),
    ("Mạng dữ liệu", "Kết nối Internet ổn định bằng giải pháp được TTTM chấp thuận để phục vụ thanh toán và đồng bộ dữ liệu."),
    ("Mặt bằng", "Nền phẳng, khô ráo, chịu tải phù hợp; không chắn lối thoát hiểm, tủ điện, họng chữa cháy hoặc tuyến giao thông chính."),
    ("Khoảng thao tác", "Bố trí đủ vùng ngồi chơi, nhận quà, nạp tiền và tiếp cận bảo trì; bảo đảm lối đi theo yêu cầu của TTTM."),
    ("Âm thanh", "Cài đặt mức âm lượng theo giới hạn của khu vực; đo và điều chỉnh tại hiện trường nếu Ban Quản lý yêu cầu."),
], widths=[4.4, 11.3], font_size=11.2)

doc.add_page_break()

add_heading("8. Nguyên lý hoạt động và hành trình người chơi", 1)
steps = [
    ("Bước 1", "Người chơi thanh toán bằng tiền mặt hoặc quét mã QR để nạp giá trị/lượt chơi."),
    ("Bước 2", "Người chơi chọn gói hoặc điểm chơi theo cấu hình được chủ máy thiết lập."),
    ("Bước 3", "Người chơi nhấn nút bắt đầu và sử dụng cần bắn bi."),
    ("Bước 4", "Bi đi qua bàn chốt; hệ thống ghi nhận vị trí rơi, hệ số thưởng và tổng điểm."),
    ("Bước 5", "Khi đủ điểm, người chơi chọn ngăn quà có mức điểm tương ứng; hệ thống mở ngăn để nhận quà."),
]
add_table(["Trình tự", "Nội dung"], steps, widths=[3.0, 12.7], font_size=11.5)

add_heading("9. Cơ chế cài đặt và kiểm soát quà", 1)
add_bullet("Chủ máy được thiết lập giá chơi, gói chơi, cách tính điểm và mức điểm đổi quà trên phần mềm.")
add_bullet("Mức điểm được hiển thị tại từng ngăn quà để người chơi biết trước điều kiện nhận quà.")
add_bullet("Hệ thống ghi nhận số lượng quà xuất ra, hỗ trợ đối chiếu doanh thu, lượt chơi và chi phí quà.")
add_bullet("Quà tặng phải phù hợp kích thước ngăn, có nguồn gốc rõ ràng và đáp ứng quy định về hàng hóa lưu thông.")
add_bullet("Điểm và quà được cấu hình theo chính sách vận hành; không sử dụng tiền mặt làm phần thưởng.")

add_heading("10. Phương thức thanh toán", 1)
add_body("Máy hỗ trợ thanh toán tiền mặt và chuyển khoản bằng mã QR. Giao dịch QR được đồng bộ về hệ thống quản lý; giao dịch tiền mặt được máy ghi nhận theo cấu hình và dữ liệu vận hành. Đơn vị vận hành thực hiện đối soát định kỳ theo từng máy và từng điểm đặt.")

doc.add_page_break()

add_heading("11. Hệ thống phần mềm quản lý", 1)
add_table(["Nhóm chức năng", "Nội dung"], [
    ("Quản lý doanh thu", "Theo dõi doanh thu theo máy, địa điểm, ngày, tháng hoặc khoảng thời gian; phân tách theo phương thức thanh toán."),
    ("Quản lý chi phí và lợi nhuận", "Ghi nhận chi phí mặt bằng, quà tặng, điện, nhân sự và khoản vận hành; hỗ trợ tổng hợp lãi/lỗ."),
    ("Quản lý quà tặng", "Theo dõi hàng hóa, tồn kho, số lượng quà đã xuất và phục vụ đối chiếu."),
    ("Quản lý giao dịch", "Theo dõi lịch sử thanh toán, nạp tiền, lượt chơi và giao dịch phát sinh."),
    ("Quản lý nhiều điểm", "Theo dõi nhiều máy, nhiều địa điểm trên cùng một hệ thống."),
    ("Cài đặt vận hành", "Thiết lập giá chơi, gói nạp, số lượt, cách tính điểm và mức điểm đổi quà."),
    ("Điều khiển từ xa", "Hỗ trợ khởi động máy và các thao tác điều khiển đã được tích hợp."),
    ("Báo cáo", "Xuất báo cáo doanh thu, chi phí, hàng hóa và biểu đồ phân tích hiệu quả."),
], widths=[4.4, 11.3], font_size=11.0)

add_heading("12. Quyền truy cập và quản trị dữ liệu", 1)
add_bullet("Tài khoản quản trị do đơn vị vận hành quản lý và phân quyền cho người có trách nhiệm.")
add_bullet("Dữ liệu doanh thu và giao dịch được đối soát theo kỳ vận hành.")
add_bullet("Thông tin kết nối mạng và thanh toán phải tuân thủ quy định an toàn thông tin, an ninh mạng và yêu cầu của TTTM.")
add_bullet("Nhà cung cấp hỗ trợ xử lý sự cố phần mềm từ xa; trường hợp không khắc phục được sẽ phối hợp kiểm tra tại địa điểm.")

doc.add_page_break()

add_heading("13. Phương án lắp đặt và vận hành tại trung tâm thương mại", 1)
add_table(["Giai đoạn", "Công việc", "Đơn vị thực hiện"], [
    ("1. Khảo sát", "Xác nhận vị trí, nguồn điện, kết nối mạng, lối đi, vùng thao tác và yêu cầu PCCC.", "Nhà cung cấp và TTTM"),
    ("2. Phê duyệt", "Chốt số lượng máy, sơ đồ bố trí, phương án nhận diện, thời gian thi công và danh mục hồ sơ.", "TTTM / đơn vị vận hành"),
    ("3. Vận chuyển", "Đưa máy đến điểm tập kết theo khung giờ và tuyến vận chuyển được TTTM chấp thuận.", "Nhà cung cấp"),
    ("4. Lắp đặt", "Đặt máy, cân chỉnh, kết nối điện/mạng, kiểm tra dây dẫn và vùng tiếp cận.", "Nhà cung cấp"),
    ("5. Kiểm thử", "Kiểm tra thanh toán, lượt chơi, ghi điểm, mở ngăn quà, phần mềm và báo cáo.", "Hai bên"),
    ("6. Bàn giao", "Lập biên bản bàn giao, hướng dẫn vận hành, thông tin bảo hành và đầu mối xử lý sự cố.", "Hai bên"),
], widths=[2.7, 9.1, 3.9], font_size=10.7)

add_heading("14. Kiểm tra đầu ca và cuối ca", 1)
add_table(["Thời điểm", "Nội dung kiểm tra"], [
    ("Đầu ca", "Ngoại quan; dây điện; kết nối mạng; thanh toán; nút bấm; cần bắn; màn hình điểm; chín ngăn quà; mức quà; khu vực xung quanh."),
    ("Trong ca", "Theo dõi cảnh báo hệ thống; bổ sung quà; vệ sinh bề mặt; xử lý giao dịch hoặc lỗi vận hành theo quy trình."),
    ("Cuối ca", "Đối soát doanh thu; kiểm đếm quà; ghi nhận sự cố; vệ sinh; tắt hoặc chuyển trạng thái máy theo quy định của TTTM."),
], widths=[3.0, 12.7], font_size=11.0)

doc.add_page_break()

add_heading("15. An toàn kỹ thuật và kiểm soát rủi ro", 1)
add_table(["Rủi ro / yêu cầu", "Biện pháp kiểm soát"], [
    ("Điện và dây dẫn", "Dùng ổ cắm đạt yêu cầu, có tiếp địa và bảo vệ rò điện; đi dây gọn, có nẹp/bảo vệ; không dùng dây hở hoặc nối tạm."),
    ("Ổn định cơ học", "Đặt trên nền phẳng; khóa bánh xe hoặc chèn cố định theo thiết kế; kiểm tra rung lắc trước khi bàn giao."),
    ("Va chạm và kẹt tay", "Kiểm tra cửa ngăn quà, bản lề, cạnh mép và cần bắn; tạm dừng máy khi phát hiện hư hỏng."),
    ("Cháy và nguồn nhiệt", "Không che khe thoát nhiệt; không đặt gần nguồn nhiệt; giữ khoảng cách với thiết bị PCCC và thực hiện yêu cầu của TTTM."),
    ("Nước và vệ sinh", "Chỉ vệ sinh bằng khăn phù hợp khi đã ngắt nguồn; không phun nước hoặc dùng hóa chất gây hư hại linh kiện."),
    ("Mạng và thanh toán", "Bảo vệ tài khoản quản trị, mã QR và thông tin truy cập; đối soát giao dịch; báo ngay khi có dấu hiệu bất thường."),
    ("Nội dung và phần thưởng", "Duy trì nội dung vui chơi lành mạnh; công khai mức điểm; chỉ đổi quà hiện vật; không quy đổi tiền, không đặt cược."),
    ("Sự cố", "Ngắt nguồn, treo thông báo ngừng phục vụ, cô lập khu vực nếu cần và liên hệ đầu mối kỹ thuật."),
], widths=[4.3, 11.4], font_size=10.8)

add_heading("16. Bảo hành và hỗ trợ kỹ thuật", 1)
add_table(["Nội dung", "Chính sách"], [
    ("Thời gian bảo hành", "12 tháng"),
    ("Phạm vi", "Phần mềm quản lý và phần cứng phát sinh lỗi do nhà cung cấp"),
    ("Hỗ trợ từ xa", "Có"),
    ("Hỗ trợ tại địa điểm", "Thực hiện khi không thể xử lý từ xa"),
    ("Thời gian phản hồi", "Trong vòng 24 giờ kể từ khi nhận thông báo"),
    ("Linh kiện", "Có chính sách thay thế; chi phí ngoài bảo hành xác định theo từng trường hợp"),
    ("Đào tạo", "Có hướng dẫn sử dụng và vận hành chi tiết"),
], widths=[4.5, 11.2], font_size=11.0)

doc.add_page_break()

add_heading("17. Hồ sơ nguồn gốc và tuân thủ", 1)
add_body("Theo tài liệu của nhà cung cấp, máy có xuất xứ Trung Quốc, được nhập khẩu chính ngạch và có hóa đơn, chứng từ mua bán. Hồ sơ kiểm tra/thẩm định nội dung và tem kiểm soát của cơ quan văn hóa có thẩm quyền được cung cấp theo từng lô hoặc từng máy, tùy chứng từ thực tế.")
add_body("Trước khi đưa máy vào trung tâm thương mại, đơn vị cung cấp và đơn vị vận hành có trách nhiệm đối chiếu model, số seri, số lượng máy, tem nhãn, tờ khai nhập khẩu, hóa đơn và các tài liệu kỹ thuật với thiết bị thực tế.")

add_heading("18. Danh mục chứng từ đề nghị đính kèm khi nộp", 1)
add_table(["STT", "Tài liệu", "Tình trạng / số tham chiếu"], [
    ("1", "Giấy chứng nhận đăng ký doanh nghiệp của đơn vị cung cấp", "Đính kèm bản sao"),
    ("2", "Hóa đơn, chứng từ mua bán", "Đính kèm theo lô/máy"),
    ("3", "Tờ khai hải quan hoặc hồ sơ nhập khẩu", "Điền số tờ khai: ………………………"),
    ("4", "Tem/nhãn kiểm soát hoặc hồ sơ thẩm định nội dung của cơ quan văn hóa", "Điền số, ngày cấp: ………………………"),
    ("5", "Catalogue hoặc tài liệu kỹ thuật của nhà sản xuất", "Đính kèm"),
    ("6", "Danh sách model, số seri và số lượng máy đưa vào TTTM", "Lập theo biên bản bàn giao"),
    ("7", "Sơ đồ vị trí lắp đặt và cấp điện", "Bổ sung sau khảo sát mặt bằng"),
    ("8", "Biên bản kiểm tra điện, chạy thử và bàn giao", "Lập tại thời điểm lắp đặt"),
    ("9", "Quy trình vận hành, xử lý sự cố và liên hệ khẩn cấp", "Đính kèm/bàn giao cho điểm đặt"),
    ("10", "Tài liệu về vật liệu, an toàn điện, EMC hoặc chứng nhận khác", "Cung cấp khi TTTM/cơ quan quản lý yêu cầu"),
], widths=[1.2, 9.2, 5.3], font_size=10.5)

add_heading("19. Đầu mối phối hợp", 1)
add_table(["Nội dung", "Thông tin"], [
    ("Đơn vị", "Công ty TNHH Công nghệ Việt Core Việt Nam"),
    ("Hotline", "037 599 3345"),
    ("Email", "tnhhcongnghevietcore@gmail.com"),
    ("Địa chỉ", "24 đường số 11, CityLand ParkHill, Gò Vấp, TP. Hồ Chí Minh"),
], widths=[4.3, 11.4], font_size=11.5)

# Images section
doc.add_page_break()
add_heading("20. Hình ảnh mô phỏng và kích thước thiết bị", 1)
add_body("Hình dưới thể hiện cấu hình hình học tham khảo của máy. Kích thước đối chiếu: rộng 510 mm, sâu 778 mm, cao 1.700 mm.", indent=False)
add_figure(IMAGES["dimensions"], "Hình 1. Ảnh mô phỏng thiết bị và các kích thước chính", width_cm=11.5)

doc.add_page_break()
add_heading("21. Hình ảnh thực tế tổng thể", 1)
add_figure(IMAGES["front"], "Hình 2. Hai máy Pinball thực tế với chín ngăn quà, bảng điểm và bàn bắn bi", width_cm=15.5)
add_body("Ảnh thực tế cho thấy hai nhóm quà mẫu gồm thú bông và sản phẩm/phụ kiện. Mức điểm đổi quà được hiển thị tại từng ngăn.", indent=False)

doc.add_page_break()
add_heading("22. Hình ảnh chi tiết bàn chơi", 1)
add_figure(IMAGES["playfield"], "Hình 3. Bàn chốt, khe ghi điểm, nút điều khiển và hướng dẫn chơi", width_cm=15.0)
add_body("Bàn chơi sử dụng cơ chế bắn bi qua hệ thống chốt; điểm được ghi nhận theo khe rơi và hệ số thưởng hiển thị trên máy.", indent=False)

doc.add_page_break()
add_heading("23. Hình ảnh chi tiết bộ điều khiển", 1)
add_figure(IMAGES["controls"], "Hình 4. Cần bắn bi, nút bắt đầu, bảng hướng dẫn và bề mặt điều khiển", width_cm=14.8)
add_body("Cụm điều khiển được bố trí ở mặt trước, thuận tiện cho người chơi thao tác và cho nhân viên kiểm tra ngoại quan trong ca vận hành.", indent=False)

doc.add_page_break()
add_heading("24. Hình ảnh bố trí thực tế", 1)
add_figure(IMAGES["installation"], "Hình 5. Bố trí hai máy và ghế ngồi tại điểm trưng bày", width_cm=12.5)
add_body("Khi lắp đặt tại trung tâm thương mại, sơ đồ cuối cùng cần thể hiện rõ vùng ngồi chơi, vùng nhận quà, lối tiếp cận bảo trì, tuyến cấp điện và khoảng lưu thông xung quanh.", indent=False)

add_heading("25. Kết luận và đề nghị", 1)
add_body("Máy bắn bi Pinball có kích thước gọn, công suất điện thấp, vận hành tự phục vụ và có hệ thống quản lý từ xa. Thiết bị phù hợp để xem xét bố trí tại khu vui chơi hoặc khu dịch vụ của trung tâm thương mại sau khi vị trí lắp đặt, kết nối điện - mạng, hồ sơ chứng từ và biện pháp an toàn được các bên xác nhận.")
add_body("Công ty TNHH Công nghệ Việt Core Việt Nam kính đề nghị Ban Quản lý Trung tâm thương mại xem xét hồ sơ, tổ chức khảo sát vị trí và hướng dẫn các yêu cầu bổ sung để hoàn thiện phương án lắp đặt.")

sig = doc.add_table(rows=1, cols=2)
sig.alignment = WD_TABLE_ALIGNMENT.CENTER
sig.autofit = False
remove_table_borders(sig)
sig.columns[0].width = Cm(7.8)
sig.columns[1].width = Cm(7.8)
for cell, title in zip(sig.rows[0].cells, ["ĐẠI DIỆN ĐƠN VỊ VẬN HÀNH", "ĐẠI DIỆN ĐƠN VỊ CUNG CẤP"]):
    set_cell_margins(cell, 100, 80, 100, 80)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, title, bold=True, size=12)
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p2, "(Ký, ghi rõ họ tên và đóng dấu)", italic=True, size=11)
    p3 = cell.add_paragraph("\n\n\n\n")

doc.core_properties.title = "Hồ sơ thuyết minh máy bắn bi Pinball trình trung tâm thương mại"
doc.core_properties.subject = "Thông số kỹ thuật, phương án lắp đặt, vận hành, an toàn và hình ảnh thiết bị"
doc.core_properties.author = "Công ty TNHH Công nghệ Việt Core Việt Nam"
doc.core_properties.keywords = "Pinball, máy bắn bi, trung tâm thương mại, hồ sơ kỹ thuật"
doc.save(OUT)
print(OUT)
